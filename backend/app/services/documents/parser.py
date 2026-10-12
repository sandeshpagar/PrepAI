import io
import json
import logging
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple
import pypdf
from docx import Document as DocxDocument
from pptx import Presentation as PptxPresentation

logger = logging.getLogger(__name__)

PS_OCR_SCRIPT = """
param([string]$inputDir)
[Console]::OutputEncoding = [System.Text.Encoding]::UTF8
$OutputEncoding = [System.Text.Encoding]::UTF8
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Graphics, ContentType = WindowsRuntime]
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType = WindowsRuntime]

$asTaskGeneric = [System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq "AsTask" -and $_.GetParameters().Count -eq 1 -and $_.IsGenericMethod } | Select-Object -First 1
$engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
if ($null -eq $engine) {
    $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage([Windows.Globalization.Language]::new("en-US"))
}

$files = Get-ChildItem -Path $inputDir -File | Where-Object { $_.Extension -match '\\.(png|jpg|jpeg|bmp)$' } | Sort-Object Name
$results = @{}

foreach ($f in $files) {
    try {
        $fileStream = [System.IO.File]::OpenRead($f.FullName)
        $randomAccessStream = [System.IO.WindowsRuntimeStreamExtensions]::AsRandomAccessStream($fileStream)
        
        $decoderAsync = [Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($randomAccessStream)
        $task = $asTaskGeneric.MakeGenericMethod([Windows.Graphics.Imaging.BitmapDecoder]).Invoke($null, @($decoderAsync))
        $decoder = $task.GetAwaiter().GetResult()
        
        $softwareBitmapAsync = $decoder.GetSoftwareBitmapAsync()
        $taskBmp = $asTaskGeneric.MakeGenericMethod([Windows.Graphics.Imaging.SoftwareBitmap]).Invoke($null, @($softwareBitmapAsync))
        $softwareBitmap = $taskBmp.GetAwaiter().GetResult()
        
        $ocrResultAsync = $engine.RecognizeAsync($softwareBitmap)
        $taskOcr = $asTaskGeneric.MakeGenericMethod([Windows.Media.Ocr.OcrResult]).Invoke($null, @($ocrResultAsync))
        $ocrResult = $taskOcr.GetAwaiter().GetResult()
        
        $fileStream.Close()
        if ($ocrResult.Text) {
            # Use newline-delimited lines for optimal structure
            $lines = $ocrResult.Lines | ForEach-Object { $_.Text.Trim() } | Where-Object { $_.Length -gt 0 }
            $results[$f.Name] = ($lines -join "`n")
        }
    } catch {
        # ignore individual image decode failures
    }
}

$results | ConvertTo-Json -Compress
"""


@dataclass
class ExtractedSection:
    content: str
    page_or_slide: Optional[int] = None


class DocumentParser:
    """
    Extracts text and page/slide labels from supported academic document formats.
    Includes hardware-accelerated local OCR for embedded diagrams and infographics.
    """

    @classmethod
    def parse(cls, file_bytes: bytes, file_type: str) -> tuple[List[ExtractedSection], int]:
        """
        Parses document bytes and returns (sections, page_or_slide_count).
        Raises ValueError on corrupted, empty, or unsupported documents.
        """
        if not file_bytes:
            raise ValueError("Document file is empty (0 bytes).")

        normalized_type = file_type.lower().strip().lstrip(".")
        if normalized_type == "pdf":
            return cls._parse_pdf(file_bytes)
        elif normalized_type == "docx":
            return cls._parse_docx(file_bytes)
        elif normalized_type == "pptx":
            return cls._parse_pptx(file_bytes)
        elif normalized_type in ("txt", "md"):
            return cls._parse_text(file_bytes)
        else:
            raise ValueError(f"Unsupported file format: '.{normalized_type}'. Supported: pdf, docx, pptx, txt, md.")

    @classmethod
    def _perform_batch_ocr(cls, images: List[Tuple[str, bytes]]) -> Dict[str, str]:
        """Runs native on-device Windows OCR on candidate diagram images."""
        if not images or os.name != "nt":
            return {}

        ps_exe = (
            shutil.which("powershell.exe")
            or shutil.which("powershell")
            or r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe"
        )
        if not os.path.exists(ps_exe) and not shutil.which(ps_exe):
            return {}

        try:
            with tempfile.TemporaryDirectory() as tmp_dir:
                for fname, img_bytes in images:
                    with open(os.path.join(tmp_dir, fname), "wb") as f:
                        f.write(img_bytes)

                script_path = os.path.join(tmp_dir, "ocr.ps1")
                with open(script_path, "w", encoding="utf-8") as f:
                    f.write(PS_OCR_SCRIPT)

                cmd = [
                    ps_exe,
                    "-NoProfile",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    script_path,
                    "-inputDir",
                    tmp_dir,
                ]
                res = subprocess.run(
                    cmd,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=30,
                )
                if res.returncode == 0 and res.stdout.strip():
                    return json.loads(res.stdout.strip())
        except Exception as e:
            logger.warning(f"Native on-device OCR skipped: {e}")
        return {}

    @classmethod
    def _parse_pdf(cls, file_bytes: bytes) -> tuple[List[ExtractedSection], int]:
        stream = io.BytesIO(file_bytes)
        try:
            reader = pypdf.PdfReader(stream)
        except Exception as e:
            raise ValueError(f"Failed to read PDF document: {str(e)}")

        total_pages = len(reader.pages)
        if total_pages == 0:
            raise ValueError("PDF has 0 pages.")

        sections: List[ExtractedSection] = []
        total_text_length = 0

        # 1. Identify candidate diagram images (>= 15KB) for OCR
        images_to_ocr: List[Tuple[str, bytes]] = []
        page_images_map: Dict[int, List[str]] = {}

        for page_num, page in enumerate(reader.pages, start=1):
            page_images_map[page_num] = []
            try:
                for img in page.images:
                    if len(img.data) >= 15_000:
                        ext = img.name.split(".")[-1].lower()
                        if ext in ("png", "jpg", "jpeg", "bmp"):
                            key = f"p{page_num}_{img.name}"
                            images_to_ocr.append((key, img.data))
                            page_images_map[page_num].append(key)
            except Exception:
                pass

        ocr_results = cls._perform_batch_ocr(images_to_ocr) if images_to_ocr else {}

        # 2. Assemble sections per page
        for page_num, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(ExtractedSection(content=text, page_or_slide=page_num))
                total_text_length += len(text)

            # Add any diagram/infographic OCR as distinct section for this page (ignore small logos < 100 chars)
            for key in page_images_map.get(page_num, []):
                ocr_text = ocr_results.get(key, "").strip()
                if ocr_text and len(ocr_text) >= 100:
                    labeled_ocr = f"[Diagram / Visual Content on Page {page_num}]:\n{ocr_text}"
                    sections.append(ExtractedSection(content=labeled_ocr, page_or_slide=page_num))
                    total_text_length += len(ocr_text)

        if total_text_length == 0:
            raise ValueError(
                "Scanned PDF or no selectable text found. Local OCR was unable to detect readable text."
            )

        return sections, total_pages

    @classmethod
    def _parse_docx(cls, file_bytes: bytes) -> tuple[List[ExtractedSection], int]:
        stream = io.BytesIO(file_bytes)
        try:
            doc = DocxDocument(stream)
        except Exception as e:
            raise ValueError(f"Failed to read DOCX document: {str(e)}")

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("DOCX document contains no extractable text.")

        # In docx, approximate pages by ~3000 chars per page
        approx_pages = max(1, len(full_text) // 3000)
        return [ExtractedSection(content=full_text, page_or_slide=1)], approx_pages

    @classmethod
    def _extract_shape_text(cls, shape) -> List[str]:
        texts: List[str] = []
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                txt = paragraph.text.strip()
                if txt:
                    texts.append(txt)
        if shape.has_table:
            for row in shape.table.rows:
                row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_txt:
                    texts.append(row_txt)
        # Recurse on grouped shapes
        if hasattr(shape, "shapes"):
            for sub_shape in shape.shapes:
                texts.extend(cls._extract_shape_text(sub_shape))
        return texts

    @classmethod
    def _parse_pptx(cls, file_bytes: bytes) -> tuple[List[ExtractedSection], int]:
        stream = io.BytesIO(file_bytes)
        try:
            prs = PptxPresentation(stream)
        except Exception as e:
            raise ValueError(f"Failed to read PPTX presentation: {str(e)}")

        total_slides = len(prs.slides)
        if total_slides == 0:
            raise ValueError("PPTX presentation has 0 slides.")

        sections: List[ExtractedSection] = []
        total_text_length = 0

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_texts: List[str] = []
            for shape in slide.shapes:
                slide_texts.extend(cls._extract_shape_text(shape))

            # Check notes slide
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"[Presenter Notes]: {notes}")

            combined = "\n".join(slide_texts).strip()
            if combined:
                sections.append(ExtractedSection(content=combined, page_or_slide=slide_num))
                total_text_length += len(combined)

        if total_text_length == 0:
            raise ValueError("PPTX slides contain no extractable text.")

        return sections, total_slides

    @classmethod
    def _parse_text(cls, file_bytes: bytes) -> tuple[List[ExtractedSection], int]:
        try:
            text = file_bytes.decode("utf-8").strip()
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1").strip()
            except Exception as e:
                raise ValueError(f"Unable to decode text document: {str(e)}")

        if not text:
            raise ValueError("Text document is empty.")

        approx_pages = max(1, len(text) // 3000)
        return [ExtractedSection(content=text, page_or_slide=1)], approx_pages
